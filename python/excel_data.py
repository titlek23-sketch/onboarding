"""원천데이터 폴더의 인사팀 엑셀 3개를 읽고 점검한다.

엑셀은 읽기 전용으로만 연다. 이 모듈은 점검(check_excel.py)과
나중에 Supabase에 올리는 프로그램이 함께 쓴다.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE_DIR = ROOT / "원천데이터"

FAQ_FILE = "01_FAQ.xlsx"
TASK_FILE = "02_할일템플릿.xlsx"
PEOPLE_FILE = "03_입사자명단.xlsx"

# (파일, 시트, 열 목록, 꼭 있어야 할 열)
SHEETS = {
    "faq": (
        FAQ_FILE, "FAQ",
        ["번호", "분류", "질문", "이런 말로도 찾음", "한 줄 답변", "상세 설명", "관련 서식", "담당 부서", "최종 확인일"],
        ["번호", "분류", "질문", "이런 말로도 찾음", "한 줄 답변", "상세 설명", "담당 부서", "최종 확인일"],
    ),
    "tasks": (
        TASK_FILE, "할일",
        ["할일ID", "시기", "기한(입사일 기준 일수)", "할 일", "설명", "담당 부서", "필수 여부", "대상 직군"],
        ["할일ID", "시기", "기한(입사일 기준 일수)", "할 일", "설명", "담당 부서", "필수 여부", "대상 직군"],
    ),
    "hires": (
        PEOPLE_FILE, "입사자",
        ["사번", "이름", "이메일", "부서", "직군", "사업장", "입사일", "멘토"],
        ["사번", "이름", "이메일", "부서", "직군", "입사일"],
    ),
    "hr": (
        PEOPLE_FILE, "HR담당자",
        ["이메일", "이름", "역할"],
        ["이메일"],
    ),
    "checks": (
        PEOPLE_FILE, "체크기록_샘플",
        ["이메일", "할일ID", "완료 일시"],
        ["이메일", "할일ID", "완료 일시"],
    ),
}

PERIODS = ["입사 전", "첫날", "첫 주", "첫 달"]
REQUIRED_KINDS = ["필수", "선택"]
TARGET_GROUPS = ["전체", "현장직", "사무직"]
JOB_GROUPS = ["현장직", "사무직"]

EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}$")
DATE_FORMATS = ["%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"]
DATETIME_FORMATS = ["%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y.%m.%d %H:%M"]


@dataclass
class Problem:
    file: str
    sheet: str
    row: int  # 엑셀 화면에 보이는 줄 번호 (제목 줄이 1)
    message: str

    def __str__(self):
        where = f"{self.file} / '{self.sheet}' 시트"
        if self.row:
            where += f" / {self.row}번째 줄"
        return f"{where}: {self.message}"


@dataclass
class ExcelData:
    faq: list = field(default_factory=list)
    tasks: list = field(default_factory=list)
    hires: list = field(default_factory=list)
    hr: list = field(default_factory=list)
    checks: list = field(default_factory=list)
    problems: list = field(default_factory=list)


def _is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == "")


def _clean(value):
    if isinstance(value, str):
        return value.strip()
    return value


def parse_date(value):
    """엑셀 날짜 칸이나 '2026-08-31' 같은 글자를 date로 바꾼다. 못 바꾸면 None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def parse_datetime(value):
    """엑셀 날짜 칸이나 '2026-08-28 14:58' 같은 글자를 datetime으로 바꾼다. 못 바꾸면 None."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str):
        text = value.strip()
        for fmt in DATETIME_FORMATS + DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
    return None


def parse_int(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"-?\d+", value.strip()):
        return int(value.strip())
    return None


def _read_sheet(key, problems):
    """시트를 읽어 [(엑셀 줄 번호, {열: 값})] 목록으로 돌려준다. 완전히 빈 줄은 건너뛴다."""
    file, sheet, columns, required = SHEETS[key]
    path = SOURCE_DIR / file
    if not path.exists():
        problems.append(Problem(file, sheet, 0, "파일이 없습니다"))
        return []

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            problems.append(Problem(file, sheet, 0, "시트가 없습니다"))
            return []
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        problems.append(Problem(file, sheet, 1, "제목 줄이 없습니다"))
        return []

    header = [_clean(h) for h in rows[0]]
    index = {}
    for col in columns:
        if col in header:
            index[col] = header.index(col)
        else:
            problems.append(Problem(file, sheet, 1, f"'{col}' 열이 없습니다"))
    if len(index) != len(columns):
        return []

    records = []
    for row_no, row in enumerate(rows[1:], start=2):
        if all(_is_blank(v) for v in row):
            continue
        record = {col: _clean(row[i]) if i < len(row) else None for col, i in index.items()}
        for col in required:
            if _is_blank(record[col]):
                problems.append(Problem(file, sheet, row_no, f"'{col}' 칸이 비어 있습니다"))
        records.append((row_no, record))
    return records


def _check_duplicates(key, records, make_key, label, problems):
    file, sheet, _, _ = SHEETS[key]
    seen = {}
    for row_no, rec in records:
        k = make_key(rec)
        if k is None or any(_is_blank(part) for part in (k if isinstance(k, tuple) else (k,))):
            continue
        if k in seen:
            problems.append(Problem(file, sheet, row_no, f"{label} '{_show(k)}'이(가) {seen[k]}번째 줄과 겹칩니다"))
        else:
            seen[k] = row_no


def _show(k):
    return " + ".join(str(p) for p in k) if isinstance(k, tuple) else str(k)


def _check_email(key, records, col, problems):
    file, sheet, _, _ = SHEETS[key]
    for row_no, rec in records:
        value = rec[col]
        if _is_blank(value):
            continue
        if not isinstance(value, str) or not EMAIL_RE.match(value):
            problems.append(Problem(file, sheet, row_no, f"'{col}' 칸이 이메일 모양이 아닙니다: {value}"))


def _check_choice(key, records, col, choices, problems):
    file, sheet, _, _ = SHEETS[key]
    for row_no, rec in records:
        value = rec[col]
        if not _is_blank(value) and value not in choices:
            problems.append(Problem(file, sheet, row_no, f"'{col}' 칸은 {'/'.join(choices)} 중 하나여야 합니다: {value}"))


def _convert_date(key, records, col, parser, problems):
    """날짜 칸을 date/datetime으로 바꿔 넣는다. 날짜 모양이 아니면 문제로 적는다."""
    file, sheet, _, _ = SHEETS[key]
    for row_no, rec in records:
        value = rec[col]
        if _is_blank(value):
            continue
        parsed = parser(value)
        if parsed is None:
            problems.append(Problem(file, sheet, row_no, f"'{col}' 칸이 날짜 모양이 아닙니다: {value}"))
        rec[col] = parsed


def load_and_check():
    """엑셀 3개를 읽고 점검한다. 결과의 problems가 비어 있어야 올릴 수 있다."""
    problems = []
    faq = _read_sheet("faq", problems)
    tasks = _read_sheet("tasks", problems)
    hires = _read_sheet("hires", problems)
    hr = _read_sheet("hr", problems)
    checks = _read_sheet("checks", problems)

    # 이메일은 대소문자를 구분하지 않는다
    for records in (hires, hr, checks):
        for _, rec in records:
            if isinstance(rec["이메일"], str):
                rec["이메일"] = rec["이메일"].lower()

    # 겹치는 번호
    _check_duplicates("faq", faq, lambda r: r["번호"], "FAQ 번호", problems)
    _check_duplicates("tasks", tasks, lambda r: r["할일ID"], "할일ID", problems)
    _check_duplicates("hires", hires, lambda r: r["사번"], "사번", problems)
    _check_duplicates("hires", hires, lambda r: r["이메일"], "이메일", problems)
    _check_duplicates("hr", hr, lambda r: r["이메일"], "이메일", problems)
    _check_duplicates("checks", checks, lambda r: (r["이메일"], r["할일ID"]), "이메일+할일ID", problems)

    # 날짜 모양
    _convert_date("faq", faq, "최종 확인일", parse_date, problems)
    _convert_date("hires", hires, "입사일", parse_date, problems)
    _convert_date("checks", checks, "완료 일시", parse_datetime, problems)

    # 이메일 모양
    _check_email("hires", hires, "이메일", problems)
    _check_email("hr", hr, "이메일", problems)
    _check_email("checks", checks, "이메일", problems)

    # 정해진 값만 들어가는 칸
    _check_choice("tasks", tasks, "시기", PERIODS, problems)
    _check_choice("tasks", tasks, "필수 여부", REQUIRED_KINDS, problems)
    _check_choice("tasks", tasks, "대상 직군", TARGET_GROUPS, problems)
    _check_choice("hires", hires, "직군", JOB_GROUPS, problems)

    file, sheet, _, _ = SHEETS["tasks"]
    for row_no, rec in tasks:
        value = rec["기한(입사일 기준 일수)"]
        if _is_blank(value):
            continue
        days = parse_int(value)
        if days is None:
            problems.append(Problem(file, sheet, row_no, f"'기한(입사일 기준 일수)' 칸이 정수가 아닙니다: {value}"))
        rec["기한(입사일 기준 일수)"] = days

    # 체크 기록이 명단과 템플릿에 있는지
    hire_emails = {r["이메일"] for _, r in hires}
    task_ids = {r["할일ID"] for _, r in tasks}
    file, sheet, _, _ = SHEETS["checks"]
    for row_no, rec in checks:
        if not _is_blank(rec["이메일"]) and rec["이메일"] not in hire_emails:
            problems.append(Problem(file, sheet, row_no, f"이메일 '{rec['이메일']}'이(가) 입사자 시트에 없습니다"))
        if not _is_blank(rec["할일ID"]) and rec["할일ID"] not in task_ids:
            problems.append(Problem(file, sheet, row_no, f"할일ID '{rec['할일ID']}'이(가) 할일 템플릿에 없습니다"))

    return ExcelData(
        faq=[r for _, r in faq],
        tasks=[r for _, r in tasks],
        hires=[r for _, r in hires],
        hr=[r for _, r in hr],
        checks=[r for _, r in checks],
        problems=problems,
    )
