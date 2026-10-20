"""주간 리포트: 입사자별 진행률을 엑셀 파일로 뽑는다.

- 저장 위치: 리포트 폴더
- 파일 이름: 온보딩_주간리포트_YYYY-MM-DD.xlsx (실행한 날, 한국 시간)
- 시트 1개 '입사자별 진행률', 맨 위에 기준일
- 칸과 순서, 계산은 대시보드의 입사자별 진행률 표와 같다

실행: .venv\\Scripts\\python python\\report.py
"""

import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from progress import build_dashboard, today_kst
from supabase_rest import ROOT, Supabase, SupabaseError, env_status

sys.stdout.reconfigure(encoding="utf-8")

REPORT_DIR = ROOT / "리포트"
SHEET = "입사자별 진행률"
HEADERS = ["이름", "부서", "직군", "입사 며칠째", "진행률", "기한 지난 수"]
WIDTHS = [12, 14, 10, 12, 10, 13]


def load():
    db = Supabase()
    hires = db.select("hires", "employee_no,name,email,department,job_group,hire_date")
    tasks = db.select("task_templates")
    checks = db.select("task_checks", "email,task_id")
    return hires, tasks, checks


def write_report(dashboard, today):
    REPORT_DIR.mkdir(exist_ok=True)
    path = REPORT_DIR / f"온보딩_주간리포트_{today.isoformat()}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    ws["A1"] = f"기준일: {today.isoformat()}"
    ws["A1"].font = Font(bold=True, size=12)

    header_row = 3
    band = PatternFill("solid", fgColor="0F766E")
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = band
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col - 1]

    if not dashboard["rows"]:
        ws.cell(row=header_row + 1, column=1, value="없음")
    for i, r in enumerate(dashboard["rows"], start=header_row + 1):
        ws.cell(row=i, column=1, value=r["name"])
        ws.cell(row=i, column=2, value=r["department"])
        ws.cell(row=i, column=3, value=r["job_group"])
        ws.cell(row=i, column=4, value=r["day_label"])
        pct = ws.cell(row=i, column=5, value=r["percent"])
        pct.number_format = '0"%"'
        ws.cell(row=i, column=6, value=r["overdue"])

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    wb.save(path)
    return path


def main():
    if any(state == "비어 있음" for _, state in env_status()):
        print(".env에 비어 있는 줄이 있어서 멈춥니다.")
        return 1

    today = today_kst()
    try:
        hires, tasks, checks = load()
    except SupabaseError as e:
        print(f"Supabase 오류: {e}")
        return 1

    dashboard = build_dashboard(hires, tasks, checks, today)
    path = write_report(dashboard, today)

    avg = dashboard["average_percent"]
    print(f"파일: {path.relative_to(ROOT)}")
    print(f"기준일: {today.isoformat()}")
    print(f"입사자 수: {dashboard['hire_count']}명")
    print(f"평균 진행률: {'-' if avg is None else f'{avg}%'}")
    print(f"기한 지난 할 일 수: {dashboard['overdue_total']}개")
    return 0


if __name__ == "__main__":
    sys.exit(main())
